# -*- coding: utf-8 -*-
"""Build the bilingual BOM / budget workbook."""
import os, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.workbook.properties import CalcProperties
import design_data as D

BASE = Path(__file__).resolve().parent.parent
OUT = BASE / "03_Spreadsheets" / "Smart_Waste_Sorting_Bin_BOM.xlsx"

GREEN = "1F4E45"
GREEN2 = "40796A"
LIME = "D8E8B1"
ORANGE = "D9822B"
BLUE = "2E6F9E"
GREY = "F2F4F3"
SOFT = "F7F9F8"
WHITE = "FFFFFF"
RED = "B03A3A"
BORDER = "D9DEE3"

thin = Side(style="thin", color=BORDER)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
header_font = Font(name="Calibri", size=10, bold=True, color=WHITE)
body_font = Font(name="Calibri", size=9.5, color="1F2937")
small_font = Font(name="Calibri", size=8.5, color="5A6778")
title_font = Font(name="Calibri", size=18, bold=True, color=GREEN)
subtitle_font = Font(name="Calibri", size=10, color="5A6778")


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = box


def style_body(ws, r1, r2, ncols):
    for r in range(r1, r2 + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = body_font
            cell.border = box
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=SOFT)


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title_block(ws, title, subtitle, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = title_font
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=subtitle)
    c.font = subtitle_font
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30


wb = Workbook()
wb.calculation = CalcProperties(fullCalcOnLoad=True)

# ---------------- README ----------------
ws = wb.active
ws.title = "README"
ws.sheet_view.showGridLines = False
set_widths(ws, [4, 24, 90, 20])
ws.merge_cells("B2:D2")
c = ws["B2"]; c.value = "Smart Waste Sorting Bin — BOM & Cost Workbook"
c.font = Font(name="Calibri", size=22, bold=True, color=GREEN)
ws.row_dimensions[2].height = 34
ws.merge_cells("B3:D3")
c = ws["B3"]; c.value = "智能分類垃圾桶 — 物料清單與成本工作簿"
c.font = Font(name="Calibri", size=14, bold=True, color=GREEN2)
ws.merge_cells("B4:D4")
c = ws["B4"]; c.value = "Version 1.0 · 29 October 2026 · Prices in HKD · Indicative only"
c.font = subtitle_font

rows = [
    ("Purpose", "Costed bill of materials for the recommended four-stream build, with tools, pin map and cut list."),
    ("Currency", "Hong Kong dollars (HKD). Prices are indicative Hong Kong / Taobao street prices in October 2026."),
    ("Recommended total", "See the BOM sheet total. Add 15% for shipping, tax and spares."),
    ("Lean build", "≈ HK$3,200: HC-SR04 instead of 4 × VL53L1X; 12 V adapter instead of battery; reuse bins; no gas struts; no SHT31/metal kit."),
    ("Enhanced build", "≈ HK$4,800: Pi 4 2 GB; aluminium frame; spare servo/stepper; better camera; UPS."),
    ("How to edit", "Change Qty or Unit price; line totals and the total update automatically. The Budget sheet uses SUMIF on the BOM."),
    ("Safety", "Use a certified Li-ion pack with BMS, a 5 A fuse, a main switch and an e-stop. Charge outside the bin. Do not use loose cells."),
    ("Academic", "Keep receipts and a logbook. Prices change; verify before ordering."),
]
r = 6
for a, b in rows:
    ws.cell(row=r, column=2, value=a).font = Font(name="Calibri", size=10, bold=True, color=GREEN)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    c = ws.cell(row=r, column=3, value=b)
    c.font = Font(name="Calibri", size=10, color="1F2937")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 32
    r += 1

# colour legend
r += 1
ws.cell(row=r, column=2, value="Priority").font = Font(name="Calibri", size=10, bold=True, color=GREEN)
ws.cell(row=r, column=3, value="Core = needed for the four-stream sorter; Optional = evidence/reliability upgrade, can be omitted in the lean build.").font = body_font
ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True)

# ---------------- BOM ----------------
ws = wb.create_sheet("BOM")
ws.sheet_view.showGridLines = False
headers = ["ID", "Group", "Item (EN)", "Item (ZH)", "Specification", "Qty", "Unit (HKD)", "Total (HKD)", "Source", "Notes (EN)", "Notes (ZH)", "Priority"]
ncols = len(headers)
title_block(ws, "Bill of Materials — recommended build", "All prices HKD, indicative. Edit Qty/Unit; Total is a formula. Core vs Optional is shown in the last column.", ncols)
hr = 4
for i, h in enumerate(headers, 1):
    ws.cell(row=hr, column=i, value=h)
style_header(ws, hr, ncols)

r = hr + 1
for row in D.BOM:
    ws.cell(row=r, column=1, value=row[0])
    ws.cell(row=r, column=2, value=row[1])
    ws.cell(row=r, column=3, value=row[2])
    ws.cell(row=r, column=4, value=row[3])
    ws.cell(row=r, column=5, value=row[4])
    ws.cell(row=r, column=6, value=row[5])
    ws.cell(row=r, column=7, value=row[6])
    ws.cell(row=r, column=8, value=f"=F{r}*G{r}")
    ws.cell(row=r, column=9, value=row[7])
    ws.cell(row=r, column=10, value=row[8])
    ws.cell(row=r, column=11, value=row[9])
    ws.cell(row=r, column=12, value=row[10])
    r += 1
last = r - 1
style_body(ws, hr + 1, last, ncols)
for rr in range(hr + 1, last + 1):
    ws.cell(row=rr, column=6).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(row=rr, column=7).number_format = '#,##0'
    ws.cell(row=rr, column=8).number_format = '#,##0'
    ws.cell(row=rr, column=8).font = Font(name="Calibri", size=9.5, bold=True, color="1F2937")
    ws.cell(row=rr, column=12).alignment = Alignment(horizontal="center", vertical="top")

# total row
tr = last + 1
ws.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=7)
c = ws.cell(row=tr, column=1, value="TOTAL / 合計")
c.font = Font(name="Calibri", size=12, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=GREEN)
c.alignment = Alignment(horizontal="right", vertical="center")
ws.cell(row=tr, column=8, value=f"=SUM(H{hr+1}:H{last})")
ws.cell(row=tr, column=8).font = Font(name="Calibri", size=12, bold=True, color=WHITE)
ws.cell(row=tr, column=8).fill = PatternFill("solid", fgColor=GREEN)
ws.cell(row=tr, column=8).number_format = '#,##0'
ws.cell(row=tr, column=8).alignment = Alignment(horizontal="center", vertical="center")
for cc in range(1, ncols + 1):
    ws.cell(row=tr, column=cc).border = box
ws.row_dimensions[tr].height = 24

# conditional formatting priority
ws.conditional_formatting.add(f"L{hr+1}:L{last}",
    FormulaRule(formula=[f'$L{hr+1}="Optional"'], fill=PatternFill("solid", fgColor="FFF2CC")))
ws.conditional_formatting.add(f"L{hr+1}:L{last}",
    FormulaRule(formula=[f'$L{hr+1}="Core"'], fill=PatternFill("solid", fgColor="E2EFDA")))

set_widths(ws, [7, 12, 28, 34, 52, 6, 10, 10, 20, 44, 44, 9])
ws.freeze_panes = f"A{hr+1}"
ws.auto_filter.ref = f"A{hr}:L{last}"
ws.print_title_rows = f"{hr}:{hr}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------- Budget ----------------
ws = wb.create_sheet("Budget")
ws.sheet_view.showGridLines = False
set_widths(ws, [4, 30, 18, 18, 55])
ws.merge_cells("B2:E2")
c = ws["B2"]; c.value = "Budget summary / 預算摘要"; c.font = title_font
ws.merge_cells("B3:E3")
c = ws["B3"]; c.value = "Group totals use SUMIF on the BOM sheet. The total is the same as the BOM total."; c.font = subtitle_font
hr = 5
for i, h in enumerate(["Group", "Cost (HKD)", "Share", "Comment"], 2):
    ws.cell(row=hr, column=i, value=h)
for c in range(2, 6):
    cell = ws.cell(row=hr, column=c)
    cell.fill = PatternFill("solid", fgColor=GREEN)
    cell.font = header_font
    cell.border = box

groups = []
for row in D.BOM:
    if row[1] not in groups:
        groups.append(row[1])
r = hr + 1
for g in groups:
    ws.cell(row=r, column=2, value=g)
    ws.cell(row=r, column=3, value=f'=SUMIF(BOM!$B$5:$B${last},B{r},BOM!$H$5:$H${last})')
    ws.cell(row=r, column=4, value=f'=C{r}/$C${hr+len(groups)+1}')
    ws.cell(row=r, column=5, value="")
    r += 1
total_row = r
ws.cell(row=total_row, column=2, value="TOTAL")
ws.cell(row=total_row, column=3, value=f'=SUM(C{hr+1}:C{total_row-1})')
ws.cell(row=total_row, column=4, value=1)
ws.cell(row=total_row, column=5, value="Recommended build, before 15% contingency.")
for rr in range(hr + 1, total_row + 1):
    for cc in range(2, 6):
        cell = ws.cell(row=rr, column=cc)
        cell.border = box
        cell.font = body_font
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if rr == total_row:
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws.cell(row=rr, column=3).number_format = '#,##0'
    ws.cell(row=rr, column=4).number_format = '0%'

# tiers
r = total_row + 3
ws.cell(row=r, column=2, value="Build tier").font = Font(name="Calibri", size=10, bold=True, color=GREEN)
ws.cell(row=r, column=3, value="Cost (HKD)").font = Font(name="Calibri", size=10, bold=True, color=GREEN)
ws.cell(row=r, column=4, value="What changes").font = Font(name="Calibri", size=10, bold=True, color=GREEN)
ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
for c in range(2, 6):
    ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=GREEN)
    ws.cell(row=r, column=c).font = header_font
    ws.cell(row=r, column=c).border = box
tiers = [
    ("Lean", 3200, "HC-SR04; 12 V adapter; reuse bins; plywood offcuts; no gas struts; no SHT31/metal kit."),
    ("Recommended", f"=C{total_row}", "As listed in the BOM; battery + adapter; 4 × VL53L1X; metal kit; gas struts."),
    ("Enhanced", 4800, "Pi 4 2 GB; aluminium frame; spare servo/stepper; better camera; UPS."),
]
for i, (a, b, c) in enumerate(tiers):
    rr = r + 1 + i
    ws.cell(row=rr, column=2, value=a).font = body_font
    ws.cell(row=rr, column=3, value=b).font = body_font
    ws.cell(row=rr, column=3).number_format = '#,##0'
    ws.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=5)
    for cc in range(2, 6):
        cell = ws.cell(row=rr, column=cc)
        cell.border = box
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.font = body_font
    ws.cell(row=rr, column=4, value=c)
    ws.cell(row=rr, column=4).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[rr].height = 42

# contingency note
r = r + 5
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
c = ws.cell(row=r, column=2, value="Recommended practice: add 15% contingency for shipping, tax, broken parts and one spare sensor. Keep every receipt for the final report.")
c.font = Font(name="Calibri", size=9, italic=True, color=RED)
c.alignment = Alignment(wrap_text=True)
ws.row_dimensions[r].height = 32

# ---------------- Tools ----------------
ws = wb.create_sheet("Tools")
ws.sheet_view.showGridLines = False
headers = ["ID", "Tool (EN)", "Tool (ZH)", "Qty", "Unit (HKD)", "Total (HKD)"]
ncols = len(headers)
title_block(ws, "Tools and workshop equipment", "If the lab already has these, budget only for consumables. The fire bag and safety glasses are not optional.", ncols)
hr = 4
for i, h in enumerate(headers, 1):
    ws.cell(row=hr, column=i, value=h)
style_header(ws, hr, ncols)
r = hr + 1
for row in D.TOOLS:
    ws.cell(row=r, column=1, value=row[0])
    ws.cell(row=r, column=2, value=row[1])
    ws.cell(row=r, column=3, value=row[2])
    ws.cell(row=r, column=4, value=row[3])
    ws.cell(row=r, column=5, value=row[4])
    ws.cell(row=r, column=6, value=f"=D{r}*E{r}")
    r += 1
last_t = r - 1
style_body(ws, hr + 1, last_t, ncols)
for rr in range(hr + 1, last_t + 1):
    ws.cell(row=rr, column=4).alignment = Alignment(horizontal="center")
    ws.cell(row=rr, column=5).number_format = '#,##0'
    ws.cell(row=rr, column=6).number_format = '#,##0'
tr = last_t + 1
ws.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=5)
c = ws.cell(row=tr, column=1, value="TOTAL TOOLS / 工具合計")
c.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=GREEN)
c.alignment = Alignment(horizontal="right")
ws.cell(row=tr, column=6, value=f"=SUM(F{hr+1}:F{last_t})")
ws.cell(row=tr, column=6).font = Font(name="Calibri", size=11, bold=True, color=WHITE)
ws.cell(row=tr, column=6).fill = PatternFill("solid", fgColor=GREEN)
ws.cell(row=tr, column=6).number_format = '#,##0'
set_widths(ws, [7, 34, 34, 7, 12, 13])
ws.freeze_panes = f"A{hr+1}"

# ---------------- Pins ----------------
ws = wb.create_sheet("Pins")
ws.sheet_view.showGridLines = False
headers = ["Signal", "ESP32 pin", "Type", "Device / bus", "Notes"]
ncols = len(headers)
title_block(ws, "ESP32 pin map", "Authoritative pin schedule for the recommended build. GPIO34-39 are input-only.", ncols)
hr = 4
for i, h in enumerate(headers, 1):
    ws.cell(row=hr, column=i, value=h)
style_header(ws, hr, ncols)
r = hr + 1
for row in D.PIN_MAP:
    for i, val in enumerate(row, 1):
        ws.cell(row=r, column=i, value=val)
    r += 1
style_body(ws, hr + 1, r - 1, ncols)
set_widths(ws, [24, 14, 20, 42, 60])
ws.freeze_panes = f"A{hr+1}"
ws.auto_filter.ref = f"A{hr}:E{r-1}"

# ---------------- Cut list ----------------
ws = wb.create_sheet("Cut list")
ws.sheet_view.showGridLines = False
headers = ["ID", "Part", "Material", "Size (mm)", "Qty", "Notes"]
ncols = len(headers)
title_block(ws, "Cut list", "Nominal dimensions. Add 1-2 mm clearance; sand to fit. Seal all plywood edges.", ncols)
hr = 4
for i, h in enumerate(headers, 1):
    ws.cell(row=hr, column=i, value=h)
style_header(ws, hr, ncols)
r = hr + 1
for row in D.CUT_LIST:
    for i, val in enumerate(row, 1):
        ws.cell(row=r, column=i, value=val)
    r += 1
style_body(ws, hr + 1, r - 1, ncols)
set_widths(ws, [7, 28, 24, 20, 7, 55])
ws.freeze_panes = f"A{hr+1}"

# ---------------- Schedule ----------------
ws = wb.create_sheet("Schedule")
ws.sheet_view.showGridLines = False
headers = ["Weeks", "Dates", "Task", "Output / milestone"]
ncols = len(headers)
title_block(ws, "24-week build schedule", "Adjust to your actual semester dates. Keep 1-week buffers after assembly, integration and before submission.", ncols)
hr = 4
for i, h in enumerate(headers, 1):
    ws.cell(row=hr, column=i, value=h)
style_header(ws, hr, ncols)
sched = [
    ("1-2", "2-15 Nov 2026", "Finalise design, BOM, order long-lead parts, proposal slides", "Proposal ready"),
    ("3-4", "16-29 Nov", "Cut frame/drawer, build bins, bench electronics", "Frame and drawer"),
    ("5-6", "30 Nov-13 Dec", "Hopper, door, load cells, spout, bearing", "First mechanical assembly"),
    ("7-8", "14-27 Dec", "ESP32 firmware, sensors, stepper/servo, serial protocol", "Subsystems work on bench"),
    ("9-10", "28 Dec-10 Jan 2027", "Dataset collection, model training, interim report", "Interim report 1st submission"),
    ("11-12", "11-24 Jan", "Integrate vision, first end-to-end sort", "First automatic sort"),
    ("13-14", "25 Jan-7 Feb", "Calibration, 4-class test, fix jams", "50% target accuracy"),
    ("15-16", "8-21 Feb", "PWA, bin levels, cloud, progress review prep", "App dashboard"),
    ("17-18", "22 Feb-7 Mar", "50-item test, reliability, safety, video", "Test report v1"),
    ("19-20", "8-21 Mar", "Refinement, enclosure, labels, user test", "Demo-ready"),
    ("21-22", "22 Mar-4 Apr", "Final testing, report draft, poster", "Report draft"),
    ("23-24", "5-18 Apr", "Final report, video, presentation, buffer", "Submission ready"),
    ("-", "27 Apr 2027", "Final deadline", "Submit"),
]
r = hr + 1
for row in sched:
    for i, val in enumerate(row, 1):
        ws.cell(row=r, column=i, value=val)
    r += 1
style_body(ws, hr + 1, r - 1, ncols)
set_widths(ws, [8, 22, 70, 26])
ws.freeze_panes = f"A{hr+1}"

# page setup for all sheets
for s in wb.worksheets:
    s.page_setup.orientation = "landscape"
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("saved", OUT)
print("sheets", wb.sheetnames)
print("bom total expected", sum(r[5] * r[6] for r in D.BOM))
