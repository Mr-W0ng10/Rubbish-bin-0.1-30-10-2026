# -*- coding: utf-8 -*-
"""Build the English and Traditional Chinese DOCX design packages."""
import os, re, sys
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Mm, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_LINE_SPACING
import design_data as D

BASE = Path(__file__).resolve().parent.parent
SRC = BASE / "04_Source"
DRAW = BASE / "02_Drawings"
OUT = BASE / "01_Documents"
OUT.mkdir(exist_ok=True)

GREEN = RGBColor(0x1F, 0x4E, 0x45)
GREEN2 = RGBColor(0x40, 0x79, 0x6A)
LIME = RGBColor(0xA3, 0xC9, 0x4A)
ORANGE = RGBColor(0xD9, 0x82, 0x2B)
BLUE = RGBColor(0x2E, 0x6F, 0x9E)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5A, 0x67, 0x78)
RED = RGBColor(0xB0, 0x3A, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY_BG = "F2F4F3"
GREEN_BG = "1F4E45"
LIME_BG = "D8E8B1"
SOFT_BG = "F7F9F8"
CALLOUT_BG = "F0F8F1"
BORDER = "D9DEE3"

BODY_FONT = "Calibri"
EA_FONT = "Microsoft JhengHei"
CODE_FONT = "Consolas"


# ---------------- low-level XML helpers ----------------
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_margins(cell, top=0.6, bottom=0.6, left=1.2, right=1.2):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for key, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = tcMar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tcMar.append(node)
        node.set(qn("w:w"), str(int(val * 56.7)))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_row_height(row, mm):
    trPr = row._tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(mm * 56.7)))
    trHeight.set(qn("w:hRule"), "atLeast")
    trPr.append(trHeight)


def para_border_bottom(paragraph, color="1F4E45", size=8):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def para_border_top(paragraph, color="D9DEE3", size=4):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(size))
    top.set(qn("w:space"), "2")
    top.set(qn("w:color"), color)
    pbdr.append(top)
    pPr.append(pbdr)


def para_shading(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_field_run(paragraph, field, placeholder="1"):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {field} "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = placeholder
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run._r.append(txt)
    run._r.append(fld3)


def set_run_font(run, size=10.5, bold=False, color=INK, italic=False, font=BODY_FONT, ea=EA_FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:eastAsia"), ea)


def add_rich_text(paragraph, text, size=10.5, color=INK, bold=False, font=BODY_FONT, ea=EA_FONT):
    """Parse **bold** and `code` inline."""
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            r = paragraph.add_run(part[2:-2])
            set_run_font(r, size=size, bold=True, color=color, font=font, ea=ea)
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            r = paragraph.add_run(part[1:-1])
            set_run_font(r, size=max(8, size - 1), bold=bold, color=RGBColor(0x2B, 0x45, 0x55), font=CODE_FONT, ea=CODE_FONT)
        else:
            r = paragraph.add_run(part)
            set_run_font(r, size=size, bold=bold, color=color, font=font, ea=ea)


# ---------------- Markdown parser ----------------
def clean_math(s):
    s = s.strip()
    s = s.replace("\\[", "").replace("\\]", "")
    s = re.sub(r"\\text\{([^}]*)\}", r"\1", s)
    s = s.replace("\\times", "×").replace("\\approx", "≈").replace("\\quad", "    ")
    s = s.replace("\\,", " ").replace("\\ ", " ").replace("\\circ", "°")
    s = s.replace("^3", "³").replace("^2", "²")
    s = s.replace("\\", "")
    return re.sub(r"\s+", " ", s).strip()


def parse_md(paths):
    lines = []
    for p in paths:
        lines.extend(Path(p).read_text(encoding="utf-8").splitlines())
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        # fenced code
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            i += 1
            blocks.append(("code", "\n".join(code)))
            continue
        # display math
        if line.strip().startswith("\\["):
            buf = [line.strip()]
            while "\\]" not in buf[-1] and i + 1 < len(lines):
                i += 1
                buf.append(lines[i].strip())
            blocks.append(("math", clean_math(" ".join(buf))))
            i += 1
            continue
        # markdown table
        if line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue
        # figure
        m = re.match(r"\[FIG:(.+?)\|(.+?)\|(\d+)\]", line.strip())
        if m:
            blocks.append(("figure", m.group(1), m.group(2), int(m.group(3))))
            i += 1
            continue
        # callout
        m = re.match(r"\[CALLOUT:(.+?)\|(.+)\]", line.strip())
        if m:
            blocks.append(("callout", m.group(1), m.group(2)))
            i += 1
            continue
        if line.strip() in ("[BOM]", "[PINMAP]", "[CUTLIST]", "[TOOLS]"):
            blocks.append(("directive", line.strip()[1:-1]))
            i += 1
            continue
        if line.strip() == "[PAGEBREAK]":
            blocks.append(("pagebreak",))
            i += 1
            continue
        if line.strip() == "---":
            blocks.append(("hr",))
            i += 1
            continue
        # headings
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            blocks.append((f"h{level}", m.group(2).strip()))
            i += 1
            continue
        # blockquote
        if line.startswith("> "):
            blocks.append(("quote", line[2:].strip()))
            i += 1
            continue
        # bullets and numbered
        if re.match(r"^-\s+", line):
            blocks.append(("bullet", re.sub(r"^-\s+", "", line)))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            blocks.append(("num", line.strip()))
            i += 1
            continue
        # paragraph: gather until blank/special
        buf = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt.strip():
                break
            if nxt.lstrip().startswith(("|", "#", "- ", "```", "\\[", "[FIG:", "[CALLOUT:", "[BOM]", "[PINMAP]", "[CUTLIST]", "[TOOLS]", "[PAGEBREAK]")):
                break
            if re.match(r"^\d+\.\s+", nxt):
                break
            if nxt.strip() == "---":
                break
            buf.append(nxt.strip())
            i += 1
        blocks.append(("p", " ".join(buf)))
    return blocks


# ---------------- document setup ----------------
def setup_document(doc, lang):
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(20)
    sec.bottom_margin = Mm(18)
    sec.left_margin = Mm(22)
    sec.right_margin = Mm(22)
    sec.header_distance = Mm(10)
    sec.footer_distance = Mm(10)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), BODY_FONT)
    rFonts.set(qn("w:hAnsi"), BODY_FONT)
    rFonts.set(qn("w:eastAsia"), EA_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in (
        ("Heading 1", 20, GREEN, 0, 10),
        ("Heading 2", 14, GREEN, 12, 4),
        ("Heading 3", 11.5, ORANGE, 10, 3),
    ):
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        rPr = st.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), BODY_FONT)
        rFonts.set(qn("w:hAnsi"), BODY_FONT)
        rFonts.set(qn("w:eastAsia"), EA_FONT)
    doc.styles["Heading 1"].paragraph_format.page_break_before = True

    # header / footer
    header = sec.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.tab_stops.add_tab_stop(Mm(166), WD_TAB_ALIGNMENT.RIGHT)
    left = "Smart Waste Sorting Bin — Engineering Design Package" if lang == "en" else "智能分類垃圾桶 — 工程設計方案"
    r = hp.add_run(left)
    set_run_font(r, size=8, color=MUTED)
    r = hp.add_run("\tv1.0")
    set_run_font(r, size=8, color=MUTED, bold=True)
    para_border_bottom(hp, "D9DEE3", 4)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page " if lang == "en" else "第 ")
    set_run_font(r, size=8, color=MUTED)
    add_field_run(fp, "PAGE")
    r = fp.add_run(" of " if lang == "en" else " 頁，共 ")
    set_run_font(r, size=8, color=MUTED)
    add_field_run(fp, "NUMPAGES")
    r = fp.add_run("" if lang == "en" else " 頁")
    set_run_font(r, size=8, color=MUTED)


def cover(doc, lang):
    # top band
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Mm(166)
    cell = table.cell(0, 0)
    cell.width = Mm(166)
    shade(cell, GREEN_BG)
    cell_margins(cell, 7, 7, 7, 7)
    set_row_height(table.rows[0], 42)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    if lang == "en":
        r = p.add_run("SMART WASTE\nSORTING BIN")
        set_run_font(r, size=30, bold=True, color=WHITE)
        p2 = cell.add_paragraph()
        r = p2.add_run("Engineering Design Package  ·  Version 1.0")
        set_run_font(r, size=11, color=RGBColor(0xD8, 0xE8, 0xB1), bold=True)
    else:
        r = p.add_run("智能分類垃圾桶")
        set_run_font(r, size=30, bold=True, color=WHITE)
        p2 = cell.add_paragraph()
        r = p2.add_run("工程設計方案  ·  版本 1.0")
        set_run_font(r, size=11, color=RGBColor(0xD8, 0xE8, 0xB1), bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    txt = ("A low-cost, vision-assisted, four-stream sorting bin\nfor a final-year engineering project"
           if lang == "en" else "低成本、視覺輔助的四分類智能垃圾桶\n適合工程系畢業專題")
    r = p.add_run(txt)
    set_run_font(r, size=16, bold=True, color=GREEN)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    txt = ("Four streams: plastic bottle · metal · wet · dry/other\n"
           "480 × 480 × 940 mm  ·  ≈ 42.8 L  ·  ≈ HK$4,050" if lang == "en" else
           "四類：膠樽 · 金屬 · 濕垃圾 · 乾／其他\n"
           "480 × 480 × 940 mm  ·  約 42.8 L  ·  約 HK$4,050")
    r = p.add_run(txt)
    set_run_font(r, size=10.5, color=MUTED)

    hero = DRAW / "cover_hero.png"
    if hero.exists():
        doc.add_picture(str(hero), width=Mm(160))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.space_after = Pt(2)
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run("Design overview  ·  設計概覽" if lang == "en" else "設計概覽  ·  Design overview")
        set_run_font(r, size=8, color=MUTED, italic=True)

    # metadata table
    meta = [
        ("Document", "Engineering design package / 工程設計方案"),
        ("Version / date", "1.0 / 29 October 2026"),
        ("Prepared for", "Final-year engineering project (CCIT4080-style)"),
        ("Design status", "Design baseline — build and verify before submission"),
        ("Language", "English" if lang == "en" else "繁體中文（香港）"),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for a, b in meta:
        row = t.add_row()
        row.cells[0].width = Mm(42)
        row.cells[1].width = Mm(124)
        shade(row.cells[0], "F2F4F3")
        for cell, txt, bold in ((row.cells[0], a, True), (row.cells[1], b, False)):
            cell_margins(cell, 0.8, 0.8, 1.5, 1.5)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(txt)
            set_run_font(r, size=8.5, bold=bold, color=INK)
    set_table_borders(t, BORDER, 4)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    txt = ("This is a design and build guide, not a submission-ready report. Build, test, photograph and write up your own results. "
           "Prices are indicative and must be verified. Safety notes are mandatory reading."
           if lang == "en" else
           "本文件是設計與製作指南，不是可直接提交的報告。請自行製作、測試、拍照和撰寫結果。價格僅供參考，購買前必須核實。安全事項必讀。")
    r = p.add_run(txt)
    set_run_font(r, size=8.5, color=RED, italic=True)


def design_at_a_glance(doc, lang):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.page_break_before = True
    r = p.add_run("Design at a glance  ·  設計一覽" if lang == "en" else "設計一覽  ·  Design at a glance")
    set_run_font(r, size=18, bold=True, color=GREEN)
    para_border_bottom(p, "1F4E45", 8)

    rows = [
        ("Application", "Indoor, one item at a time, four streams" if lang == "en" else "室內、每次一件、四分類"),
        ("Streams", "Plastic bottle · metal · wet · dry/other" if lang == "en" else "膠樽 · 金屬 · 濕垃圾 · 乾／其他"),
        ("Size / mass", "480 × 480 × 940 mm; ≈ 22–26 kg empty" if lang == "en" else "480 × 480 × 940 mm；空機約 22–26 kg"),
        ("Capacity", "4 × 10.7 L usable; ≈ 42.8 L total" if lang == "en" else "4 × 10.7 L 可用；合計約 42.8 L"),
        ("Cycle", "6–9 s typical; ≤ 10 s target" if lang == "en" else "一般 6–9 秒；目標 ≤10 秒"),
        ("Accuracy", "≥ 85% overall on a 50-item mixed test" if lang == "en" else "50 件混合測試整體 ≥85%"),
        ("Control", "Raspberry Pi Zero 2 W (vision/app) + ESP32 (real-time)" if lang == "en" else "Raspberry Pi Zero 2 W（視覺／App）＋ESP32（實時）"),
        ("Power", "12 V 6 Ah Li-ion (BMS) or 12 V 5 A adapter" if lang == "en" else "12 V 6 Ah 鋰電（BMS）或 12 V 5 A 適配器"),
        ("Cost", "≈ HK$4,050 recommended; ≈ HK$3,200 lean" if lang == "en" else "建議約 HK$4,050；精簡約 HK$3,200"),
        ("Safety", "Fused; e-stop; lid/drawer/door interlocks; spring-closed door" if lang == "en" else "保險絲；急停；蓋／抽屜／活門連鎖；彈簧關門"),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (a, b) in enumerate(rows):
        row = t.add_row()
        row.cells[0].width = Mm(38)
        row.cells[1].width = Mm(128)
        shade(row.cells[0], "E8F0EC")
        for cell, txt, bold in ((row.cells[0], a, True), (row.cells[1], b, False)):
            cell_margins(cell, 0.9, 0.9, 1.5, 1.5)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(txt)
            set_run_font(r, size=8.8, bold=bold, color=INK)
    set_table_borders(t, BORDER, 4)

    # three "why" boxes
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Why this design" if lang == "en" else "為何這樣設計")
    set_run_font(r, size=12, bold=True, color=GREEN)
    why = [
        ("Low torque, compact" if lang == "en" else "低扭矩、緊湊",
         "The bins stay fixed; only a 0.35 kg spout rotates. No 15 kg carousel, no 700 mm round enclosure." if lang == "en"
         else "桶固定，只有 0.35 kg 滑槽轉動。不需要轉 15 kg 轉盤，也不需要 700 mm 圓形外殼。"),
        ("Safe by default" if lang == "en" else "預設安全",
         "No actuator moves unless the lid, drawer and door are closed. The e-stop cuts actuator 12 V, not the controllers." if lang == "en"
         else "蓋、抽屜、活門未關，執行器不動作。急停切斷執行器 12 V，控制器保持供電。"),
        ("Buildable and repairable" if lang == "en" else "可製作、可維修",
         "Plywood, standard bins, 3D-printed brackets and off-the-shelf electronics. Bins and liner lift out for cleaning." if lang == "en"
         else "夾板、標準桶、3D 打印支架和現成電子零件。桶和內襯可取出清潔。"),
    ]
    t = doc.add_table(rows=1, cols=3)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (title, body) in enumerate(why):
        cell = t.cell(0, i)
        cell.width = Mm(55)
        shade(cell, "F0F8F1")
        cell_margins(cell, 2, 2, 2, 2)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        set_run_font(r, size=9.5, bold=True, color=GREEN)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        r = p2.add_run(body)
        set_run_font(r, size=8, color=INK)
    set_table_borders(t, "C9D8D0", 4)

    # Build roadmap strip
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Build roadmap" if lang == "en" else "製作路線圖")
    set_run_font(r, size=12, bold=True, color=GREEN)
    roadmap = [
        ("1", "Frame + drawer", "Weeks 3–6" if lang == "en" else "第 3–6 週", "Cut, assemble, fit bins" if lang == "en" else "切割、組裝、放桶"),
        ("2", "Hopper + spout", "Weeks 5–8" if lang == "en" else "第 5–8 週", "Load cells, door, bearing" if lang == "en" else "負重、活門、軸承"),
        ("3", "Electronics + firmware", "Weeks 7–10" if lang == "en" else "第 7–10 週", "Bench test every sensor" if lang == "en" else "逐個感測器測試"),
        ("4", "Vision + integration", "Weeks 10–14" if lang == "en" else "第 10–14 週", "Dataset, model, first sort" if lang == "en" else "資料集、模型、第一次分類"),
        ("5", "Test + report", "Weeks 15–24" if lang == "en" else "第 15–24 週", "50-item test, video, report" if lang == "en" else "50 件測試、影片、報告"),
    ]
    rt = doc.add_table(rows=0, cols=4)
    rt.autofit = False
    rt.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (n, phase, weeks, out) in enumerate(roadmap):
        row = rt.add_row()
        row.cells[0].width = Mm(8)
        row.cells[1].width = Mm(42)
        row.cells[2].width = Mm(30)
        row.cells[3].width = Mm(86)
        shade(row.cells[0], GREEN_BG)
        for ci, txt in enumerate((n, phase, weeks, out)):
            cell = row.cells[ci]
            cell_margins(cell, 0.7, 0.7, 1.2, 1.2)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if ci == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(txt)
                set_run_font(r, size=10, bold=True, color=WHITE)
            elif ci == 1:
                r = p.add_run(txt)
                set_run_font(r, size=9, bold=True, color=INK)
            else:
                r = p.add_run(txt)
                set_run_font(r, size=8.5, color=INK if ci == 3 else ORANGE, bold=(ci == 2))
    set_table_borders(rt, BORDER, 4)
    # The next H1 starts a new page; no explicit break needed.


# ---------------- block renderers ----------------
def add_heading(doc, text, level, lang):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    if level == 1:
        is_appendix = text.startswith("Appendix") or text.startswith("附錄")
        is_first_appendix = text.startswith("Appendix A") or text.startswith("附錄 A")
        # Major chapters start on a new page. Appendices B-F flow on from the
        # previous table to avoid a near-empty page after a long BOM/cut list.
        p.paragraph_format.page_break_before = (not is_appendix) or is_first_appendix
    add_rich_text(p, text, size={1: 20, 2: 14, 3: 11.5}[level],
                  color={1: GREEN, 2: GREEN, 3: ORANGE}[level], bold=True)
    if level == 1:
        para_border_bottom(p, "1F4E45", 8)
    return p


def add_paragraph(doc, text, lang):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_rich_text(p, text, size=10.5)
    return p


def add_bullet(doc, text, lang):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Mm(6)
    add_rich_text(p, text, size=10.5)
    return p


def add_num(doc, text, lang):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-6)
    add_rich_text(p, text, size=10.5)
    return p


def add_code(doc, code, lang):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.left_indent = Mm(2)
    para_shading(p, "F4F6F8")
    for i, line in enumerate(code.split("\n")):
        r = p.add_run(line)
        set_run_font(r, size=7.8, color=RGBColor(0x2B, 0x45, 0x55), font=CODE_FONT, ea=CODE_FONT)
        if i < len(code.split("\n")) - 1:
            r.add_break()
    return p


def add_math(doc, text, lang):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, size=10.5, italic=True, color=INK, font="Cambria", ea=EA_FONT)
    return p


def add_quote(doc, text, lang):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Mm(5)
    p.paragraph_format.right_indent = Mm(5)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    para_border_top(p, "A3C94A", 12)
    r = p.add_run(text)
    set_run_font(r, size=10.5, italic=True, color=GREEN)
    return p


def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    para_border_top(p, "D9DEE3", 4)
    return p


def add_figure(doc, filename, caption, width_mm, lang):
    path = DRAW / filename
    if not path.exists():
        p = doc.add_paragraph(f"[missing figure: {filename}]")
        return
    with Image.open(path) as im:
        w, h = im.size
    width_mm = min(width_mm, 166)
    height_mm = width_mm * h / w
    if height_mm > 215:
        width_mm = 215 * w / h
        height_mm = 215
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Mm(width_mm))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    add_rich_text(cp, caption, size=8.2, color=MUTED, bold=False)
    for r in cp.runs:
        r.font.italic = True
    return p


def add_callout(doc, title, body, lang):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    cell.width = Mm(166)
    shade(cell, CALLOUT_BG)
    cell_margins(cell, 2.2, 2.2, 3, 3)
    # left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), "1F4E45")
    borders.append(left)
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run_font(r, size=9.8, bold=True, color=GREEN)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    add_rich_text(p2, body, size=9.2)
    # spacing
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.space_before = Pt(0)
    for r in sp.runs:
        r.font.size = Pt(2)


def add_table(doc, rows, lang, font_size=None, widths=None):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    if font_size is None:
        font_size = 8.2 if ncols <= 7 else 7.2
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t, BORDER, 4)
    if widths is None:
        widths = [166 / ncols] * ncols
    for ri, row in enumerate(rows):
        tr = t.add_row()
        set_cant_split(tr)
        for ci, val in enumerate(row):
            cell = tr.cells[ci]
            cell.width = Mm(widths[ci] if ci < len(widths) else 166 / ncols)
            cell_margins(cell, 0.5, 0.5, 1.0, 1.0)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if ri == 0:
                shade(cell, GREEN_BG)
                add_rich_text(p, val, size=font_size, color=WHITE, bold=True)
                set_repeat_header(tr)
            else:
                if ri % 2 == 0:
                    shade(cell, SOFT_BG)
                add_rich_text(p, val, size=font_size, color=INK)
    # spacing
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.line_spacing = 1.0


def add_bom(doc, lang):
    headers = (["ID", "Item", "Specification", "Qty", "Unit (HKD)", "Total (HKD)", "Priority"] if lang == "en"
               else ["編號", "項目", "規格", "數量", "單價（HKD）", "小計（HKD）", "優先"])
    widths = [11, 40, 58, 10, 15, 15, 17]
    # count rows
    groups = []
    for row in D.BOM:
        if row[1] not in groups:
            groups.append(row[1])
    total_rows = 1 + len(groups) + len(D.BOM) + 1
    t = doc.add_table(rows=0, cols=7)
    t.style = "Table Grid"
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t, BORDER, 4)
    # header
    tr = t.add_row(); set_cant_split(tr)
    for ci, val in enumerate(headers):
        cell = tr.cells[ci]
        cell.width = Mm(widths[ci])
        cell_margins(cell, 0.5, 0.5, 0.8, 0.8)
        shade(cell, GREEN_BG)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_rich_text(p, val, size=7.8, color=WHITE, bold=True)
    set_repeat_header(tr)
    for g in groups:
        items = [r for r in D.BOM if r[1] == g]
        gtotal = sum(r[5] * r[6] for r in items)
        tr = t.add_row(); set_cant_split(tr)
        merged = tr.cells[0].merge(tr.cells[6])
        shade(merged, LIME_BG)
        merged.width = Mm(166)
        cell_margins(merged, 0.7, 0.7, 1.2, 1.2)
        p = merged.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        label = f"{g.upper()}  ·  HK${gtotal:,}" if lang == "en" else f"{g}  ·  HK${gtotal:,}"
        r = p.add_run(label)
        set_run_font(r, size=8, bold=True, color=GREEN)
        for row in items:
            tr = t.add_row(); set_cant_split(tr)
            vals = [row[0], row[2] if lang == "en" else row[3], row[4], str(row[5]),
                    f"{row[6]:,}", f"{row[5]*row[6]:,}", row[10]]
            for ci, val in enumerate(vals):
                cell = tr.cells[ci]
                cell.width = Mm(widths[ci])
                cell_margins(cell, 0.4, 0.4, 0.8, 0.8)
                if len(t.rows) % 2 == 0:
                    shade(cell, SOFT_BG)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                add_rich_text(p, val, size=7.0, color=INK)
    # total
    total = sum(r[5] * r[6] for r in D.BOM)
    tr = t.add_row(); set_cant_split(tr)
    merged = tr.cells[0].merge(tr.cells[6])
    shade(merged, GREEN_BG)
    merged.width = Mm(166)
    cell_margins(merged, 0.8, 0.8, 1.2, 1.2)
    p = merged.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"TOTAL  HK${total:,}" if lang == "en" else f"合計  HK${total:,}")
    set_run_font(r, size=9.5, bold=True, color=WHITE)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)


def add_pinmap(doc, lang):
    headers = (["Signal", "ESP32 pin", "Type", "Device / bus", "Notes"] if lang == "en"
               else ["訊號", "ESP32 接腳", "類型", "裝置／匯流排", "備註"])
    rows = [headers] + [[a, b, c, d, e] for a, b, c, d, e in D.PIN_MAP]
    add_table(doc, rows, lang, font_size=7.2, widths=[31, 20, 22, 51, 42])


def add_cutlist(doc, lang):
    headers = (["ID", "Part", "Material", "Size (mm)", "Qty", "Notes"] if lang == "en"
               else ["編號", "零件", "材料", "尺寸（mm）", "數量", "備註"])
    rows = [headers] + [[a, b, c, d, str(e), f] for a, b, c, d, e, f in D.CUT_LIST]
    add_table(doc, rows, lang, font_size=7.5, widths=[11, 40, 33, 28, 10, 44])


def add_tools(doc, lang):
    headers = (["ID", "Tool", "Description", "Qty", "Cost (HKD)"] if lang == "en"
               else ["編號", "工具", "說明", "數量", "成本（HKD）"])
    rows = [headers] + [[a, b if lang == "en" else c, c if lang == "en" else b, str(q), f"{u:,}"] for a, b, c, q, u in D.TOOLS]
    add_table(doc, rows, lang, font_size=7.8, widths=[12, 55, 55, 12, 32])


# ---------------- build ----------------
def build(lang):
    if lang == "en":
        files = [SRC / f"content_en_{i}.md" for i in range(1, 6)]
        out = OUT / "Smart_Waste_Sorting_Bin_Design_EN.docx"
    else:
        files = [SRC / f"content_zh_{i}.md" for i in range(1, 6)]
        out = OUT / "智能分類垃圾桶_設計方案_ZH-Hant.docx"
    blocks = parse_md(files)
    # The Markdown files begin with a title block already reproduced on the cover.
    # Start the body at the "How to use" heading.
    start_marker = "How to use this document" if lang == "en" else "如何使用本文件"
    for idx, b in enumerate(blocks):
        if b[0] == "h1" and b[1].strip() == start_marker:
            blocks = blocks[idx:]
            break
    doc = Document()
    setup_document(doc, lang)
    cover(doc, lang)
    design_at_a_glance(doc, lang)

    for block in blocks:
        kind = block[0]
        if kind == "h1":
            add_heading(doc, block[1], 1, lang)
        elif kind == "h2":
            add_heading(doc, block[1], 2, lang)
        elif kind == "h3":
            add_heading(doc, block[1], 3, lang)
        elif kind == "p":
            add_paragraph(doc, block[1], lang)
        elif kind == "bullet":
            add_bullet(doc, block[1], lang)
        elif kind == "num":
            add_num(doc, block[1], lang)
        elif kind == "code":
            add_code(doc, block[1], lang)
        elif kind == "math":
            add_math(doc, block[1], lang)
        elif kind == "hr":
            add_hr(doc)
        elif kind == "quote":
            add_quote(doc, block[1], lang)
        elif kind == "figure":
            add_figure(doc, block[1], block[2], block[3], lang)
        elif kind == "callout":
            add_callout(doc, block[1], block[2], lang)
        elif kind == "table":
            add_table(doc, block[1], lang)
        elif kind == "directive":
            if block[1] == "BOM":
                add_bom(doc, lang)
            elif block[1] == "PINMAP":
                add_pinmap(doc, lang)
            elif block[1] == "CUTLIST":
                add_cutlist(doc, lang)
            elif block[1] == "TOOLS":
                add_tools(doc, lang)
        elif kind == "pagebreak":
            doc.add_page_break()

    # document properties
    doc.core_properties.title = "Smart Waste Sorting Bin — Engineering Design Package"
    doc.core_properties.subject = "Low-cost four-stream smart waste bin"
    doc.core_properties.author = "Engineering design package"
    doc.core_properties.keywords = "smart bin, waste sorting, ESP32, Raspberry Pi, MobileNetV2"
    doc.save(out)
    print("saved", out)
    # verify
    d = Document(out)
    print("  paragraphs", len(d.paragraphs), "tables", len(d.tables), "sections", len(d.sections))
    return out


if __name__ == "__main__":
    build("en")
    build("zh")
